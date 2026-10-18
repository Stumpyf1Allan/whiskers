"""Tests for 1.1: the optional Mortgage and Retirement sections, the plan builder with its
theme views, parachutes and brief, importing files, and the catalogue of firms.

Grouped like the rest: by what would go wrong for a person if the test didn't exist.
"""

from __future__ import annotations

import datetime as dt
import json
import urllib.error
import urllib.request

import pytest

from folio import config, db, inbox, reminders
from folio.engine import (builder, mortgage as mort, pensions, plan, retirement, rules, sample, sections,
                          statepension, themes, uktax, xray)
from folio.importers import investments as inv

from .test_whiskers import add_platform, add_position


def _years_ago(n: float) -> str:
    return (dt.date.today() - dt.timedelta(days=round(n * 365.25))).isoformat()


def _on(**kw):
    """Switch sections on: show_mortgage=True, show_retirement=True."""
    config.settings.update({f"show_{k}": v for k, v in kw.items()})


def _basic_plan(**kw):
    s = {"dob": _years_ago(60.3), "stop_age": 60.3, "spend_essential": 1200, "spend_extra": 300,
         "horizon_age": 95, "nation": "england"}
    s.update(kw)
    retirement.save(s)


# ============================================================================ tax, the State Pension

def test_income_tax_and_national_insurance_on_50k():
    assert uktax.income_tax(50000, "england")["tax"] == pytest.approx(7486.0)
    assert uktax.national_insurance(50000) == pytest.approx(2994.40)
    assert uktax.income_tax(50000, "scotland")["tax"] == pytest.approx(8982.05)


@pytest.mark.parametrize("dob,months", [
    ("1959-12-31", 66 * 12), ("1960-04-06", 66 * 12 + 1), ("1960-05-10", 66 * 12 + 2),
    ("1961-03-06", 67 * 12), ("1977-04-05", 67 * 12), ("1978-04-06", 68 * 12),
])
def test_state_pension_age_follows_the_timetable_in_law(dob, months):
    assert statepension.age_months(dt.date.fromisoformat(dob))[0] == months


def test_a_forecast_beats_a_rough_figure_and_gaps_are_costed():
    dob = dt.date(1966, 3, 14)
    est = statepension.estimate({"state_pension": {"weekly": 221.5}}, dob)
    assert est["yearly"] == pytest.approx(221.5 * 52) and est["basis"] == "your forecast"
    rough = statepension.estimate({"state_pension": {"years": 9}, "stop_age": 60}, dob)
    assert rough["basis"].startswith("rough") and rough["weekly"] == 0.0     # under ten years: nothing
    gaps = statepension.estimate({"state_pension": {"years": 25, "gaps": 3}, "stop_age": 60}, dob)
    assert gaps["top_up"]["years"] == 3 and gaps["top_up"]["pays_back_years"] < 4


# ============================================================================ the mortgage

def _mortgage(**kw) -> int:
    row = {"name": "Home", "lender": "Nationwide", "statement_balance": 100000,
           "statement_on": dt.date.today().replace(day=1).isoformat(), "monthly_payment": 527.84, "rate": 4.0,
           "repayment_type": "repayment", "rate_type": "fixed", "created_at": db.now()}
    row.update(kw)
    with db.tx() as c:
        return c.execute(f"INSERT INTO mortgages({','.join(row)}) VALUES({','.join('?' * len(row))})",
                         tuple(row.values())).lastrowid


def test_a_payment_and_its_term_agree():
    assert mort.payment_for(100000, 4.0, 300) == pytest.approx(527.84, abs=0.01)
    assert mort.months_to_clear(100000, 4.0, 527.84) == 300


def test_a_payment_is_split_into_interest_and_capital():
    s = mort.summarise(mort.get(_mortgage()))
    assert s["interest_this_month"] == pytest.approx(333.33, abs=0.01)
    assert s["capital_this_month"] == pytest.approx(527.84 - 333.33, abs=0.02)
    assert s["months_left"] == 300


def test_overpaying_saves_interest_and_time_and_nothing_is_free():
    m = mort.get(_mortgage())
    r = mort.overpayment_effect(m, 200, "monthly")
    assert r["available"] and r["months_saved"] > 60 and r["interest_saved"] > 10000
    assert "Check yours" in r["caveat"]
    assert not mort.overpayment_effect(m, 0)["available"]


def test_a_payment_that_doesnt_cover_the_interest_is_said_so():
    s = mort.summarise(mort.get(_mortgage(monthly_payment=300)))
    assert s["never_clears"] is True


def test_an_interest_only_overpayment_cuts_what_is_left_at_the_end():
    start = dt.date.today().replace(day=1)
    mid = _mortgage(repayment_type="interest_only", statement_balance=120000, monthly_payment=500, rate=5.0,
                    started_on=start.replace(year=start.year - 5).isoformat(), term_months=120)
    m = mort.get(mid)
    r = mort.overpayment_effect(m, 200, "monthly")
    assert r["available"] and r["interest_only"]
    assert r["end_balance_after"] == pytest.approx(120000 - 200 * 60)
    assert r["interest_saved"] == pytest.approx(200 * 5 / 1200 * 60 * 61 / 2, abs=1)
    assert mort.rate_change_effect(m, 6.0)["payment_then"] == pytest.approx(600)


def test_a_rate_rise_raises_the_payment():
    r = mort.rate_change_effect(mort.get(_mortgage()), 6.0)
    assert r["available"] and r["payment_then"] > r["payment_now"]


def test_the_home_has_one_value_wherever_it_is_read():
    with db.tx() as c:
        home = c.execute("INSERT INTO assets(name,kind,value,is_home) VALUES('Home','property',300000,1)").lastrowid
    _mortgage(property_asset_id=home, property_value=999999)
    _on(mortgage=True)
    bs = sections.balance_sheet()
    assert bs["own"]["Home and property"] == 300000              # not 300,000 + 999,999
    assert mort.summarise(mort.all_mortgages()[0])["property_value"] == 300000


# ============================================================================ pensions and the retirement plan

def test_a_final_salary_pension_is_income_not_money_owned():
    pensions.add({"name": "NHS Pension", "kind": "db", "yearly": 12000, "normal_age": 67})
    _on(retirement=True)
    bs = sections.balance_sheet()
    assert "Pensions" not in bs["own"] and bs["promised"][0]["yearly"] == 12000


def test_a_pension_needs_a_name_and_sensible_numbers():
    with pytest.raises(pensions.PensionError):
        pensions.add({"kind": "dc", "pot": 1000})
    with pytest.raises(pensions.PensionError):
        pensions.add({"name": "Pot", "kind": "dc", "pot": "lots"})


def test_a_plan_needs_the_basics_before_it_runs():
    out = retirement.analyse(full=False)
    assert out["ready"] is False and "your date of birth" in out["missing"]


def test_enough_money_lasts_and_too_little_says_when_it_runs_short():
    pid = add_platform(cash=0)
    with db.tx() as c:
        c.execute("UPDATE platforms SET cash=? WHERE id=?", (900000, pid))
    _basic_plan()
    assert retirement.analyse(full=False)["plan"]["lasts"] is True
    with db.tx() as c:
        c.execute("UPDATE platforms SET cash=? WHERE id=?", (20000, pid))
    out = retirement.analyse(full=True)
    assert out["plan"]["lasts"] is False and out["plan"]["runs_out_at"] < 95
    assert out["levers"], "a plan that falls short says what would fix it"


def test_a_final_salary_pension_taken_early_is_reduced():
    _basic_plan(stop_age=62)
    pensions.add({"name": "NHS", "kind": "db", "yearly": 10000, "normal_age": 67, "start_age": 62, "early_reduction": 5})
    rows = retirement.analyse(full=False)["plan"]["rows"]
    paying = [r["db"] for r in rows if r["db"]]
    assert paying[0] == pytest.approx(7500, abs=1)               # five years early at 5% a year


def test_a_pension_already_being_paid_brings_no_second_lump_sum():
    _basic_plan(dob=_years_ago(64.5), stop_age=60)
    pensions.add({"name": "NHS 1995", "kind": "db", "yearly": 16000, "normal_age": 60, "start_age": 60,
                  "lump_sum": 48000, "early_reduction": 5})
    rows = retirement.analyse(full=False)["plan"]["rows"]
    assert all(r["lump"] == 0 for r in rows)                      # it was paid at 60
    assert rows[0]["db"] == pytest.approx(16000)                  # what it pays now, as typed


def test_a_pot_already_in_drawdown_has_had_its_tax_free_cash():
    _basic_plan(dob=_years_ago(66), stop_age=60)
    pensions.add({"name": "Old pot", "kind": "dc", "pot": 80000, "start_age": 61})
    rows = retirement.analyse(full=False)["plan"]["rows"]
    assert all(r["lump"] == 0 for r in rows)                      # had at 61
    assert rows[0]["from_pension"] > 0 or rows[0]["from_cash"] >= 0


def test_a_pension_pot_is_opened_before_the_plan_is_called_short():
    # Still working at 63 when a big bill falls due: the pot pays (with its tax), rather
    # than the plan running short with a pension untouched.
    _basic_plan(dob=_years_ago(62.6), stop_age=66, salary=40000, spend_essential=1500, spend_extra=300,
                one_offs=[{"label": "Interest-only mortgage ends", "age": 63, "amount": 60000}])
    pensions.add({"name": "Workplace pot", "kind": "dc", "pot": 200000})
    plan = retirement.analyse(full=False)["plan"]
    year = next(r for r in plan["rows"] if int(r["age"]) == 63)
    assert year["short"] == 0 and year["lump"] > 0                # tax-free cash, then drawdown
    assert plan["runs_out_at"] is None or plan["runs_out_at"] > 64


def test_the_state_pension_starts_part_way_through_its_first_year():
    _basic_plan(dob=_years_ago(60.5), stop_age=60.5)              # born in 1966: State Pension at 67
    retirement.save({"state_pension": {"weekly": 200}})
    rows = retirement.analyse(full=False)["plan"]["rows"]
    first = next(r for r in rows if r["state"])
    assert 0 < first["state"] < 200 * 52                          # the half-year after 67
    assert rows[rows.index(first) + 1]["state"] == pytest.approx(200 * 52)


def test_a_uk_state_pension_paid_in_south_africa_is_frozen():
    _basic_plan(dob=_years_ago(68), stop_age=68)
    retirement.save({"state_pension": {"weekly": 200}, "retire_country": "za"})
    rows = retirement.analyse(full=False)["plan"]["rows"]
    assert rows[5]["state"] < rows[1]["state"]                     # less each year in today's money
    retirement.save({"retire_country": "uk"})
    rows = retirement.analyse(full=False)["plan"]["rows"]
    assert rows[5]["state"] == pytest.approx(rows[1]["state"])


def test_after_a_fixed_deal_the_plan_assumes_a_new_one_not_the_standard_rate():
    _basic_plan()
    _on(mortgage=True)
    end = (dt.date.today() + dt.timedelta(days=90)).isoformat()
    _mortgage(fixed_until=end, revert_rate=9.0, rate=3.0, monthly_payment=474.21)
    s, data = retirement.settings(), retirement.gather()
    base = retirement.project(s, data)["rows"]
    svr = retirement.project(s, data, {"revert_rate": "svr"})["rows"]
    assert base[2]["mortgage"] < svr[2]["mortgage"]


def test_a_fixed_payment_weighs_less_each_year_in_todays_money():
    _basic_plan()
    _mortgage()
    _on(mortgage=True)
    rows = retirement.analyse(full=False)["plan"]["rows"]
    assert rows[1]["mortgage"] > rows[5]["mortgage"] > 0


def test_pensions_left_over_count_for_inheritance_tax_unless_left_to_a_spouse():
    pid = add_platform(cash=0)
    with db.tx() as c:
        c.execute("UPDATE platforms SET cash=? WHERE id=?", (2000000, pid))
    _basic_plan()
    out = retirement.analyse(full=False)["plan"]
    assert out["iht"]["tax"] > 0
    retirement.save({"partner": True})
    assert retirement.analyse(full=False)["plan"]["iht"]["tax"] == 0


def test_a_what_if_saves_nothing():
    _basic_plan()
    before = retirement.settings()
    retirement.project(before, retirement.gather(), {"stop_age": 70})
    assert retirement.settings()["stop_age"] == before["stop_age"]


# ============================================================================ words

def test_titles_keep_their_capitals_inside_a_sentence():
    assert themes.lower_first("The AI trade") == "the AI trade"
    assert themes.lower_first("UK companies") == "UK companies"
    assert themes.lower_first("Gold") == "gold"


# ============================================================================ the X-Ray

def test_the_xray_sees_through_a_world_fund():
    pid = add_platform()
    add_position(pid, "VWRP", "LSE", qty=100, broker_value=10000.0, valued_at=db.now())
    x = xray.run()
    ai = next(t for t in x["themes"] if t["theme"] == "ai")
    assert ai["pct"] == pytest.approx(0.206, abs=0.01)
    assert any(b["company"] for b in x["biggest_companies"])


# ============================================================================ files in, the folder, the reminder

def test_a_holdings_file_is_shown_before_it_goes_in_and_replaces_the_last_one():
    pid = add_platform("Hargreaves Lansdown ISA", "other")
    rep = inv.analyse(inv.sample_holdings_csv(), "hl-holdings.csv")
    assert rep["kind"] == "holdings" and rep["rows"] == 3 and rep["suggested_platform"] == pid
    assert not db.scalar("SELECT COUNT(*) FROM positions", (), 0)          # analysing writes nothing
    inv.commit(rep["token"], pid)
    assert db.scalar("SELECT COUNT(*) FROM positions WHERE platform_id=?", (pid,)) == 3
    rep2 = inv.analyse(inv.sample_holdings_csv(), "hl-holdings.csv")
    inv.commit(rep2["token"], pid)
    assert db.scalar("SELECT COUNT(*) FROM positions WHERE platform_id=?", (pid,)) == 3


def test_the_statements_folder_imports_what_it_can_place_and_leaves_the_rest(tmp_path):
    folder = tmp_path / "statements"
    config.settings.update({"statements_folder": str(folder)})
    pid = add_platform("Hargreaves Lansdown ISA", "other")
    inbox.folder()
    (folder / "hargreaves holdings.csv").write_bytes(inv.sample_holdings_csv())
    (folder / "mystery.csv").write_text("a,b\n1,2\n", encoding="utf-8")
    out = inbox.import_waiting()
    assert [d["file"] for d in out["done"]] == ["hargreaves holdings.csv"]
    assert [x["file"] for x in out["left"]] == ["mystery.csv"]
    assert (folder / "mystery.csv").exists() and not (folder / "hargreaves holdings.csv").exists()
    assert db.scalar("SELECT COUNT(*) FROM positions WHERE platform_id=?", (pid,)) == 3
    # the same file again, renamed: recognised by its contents and filed, not imported twice
    (folder / "copy.csv").write_bytes(inv.sample_holdings_csv())
    again = inbox.import_waiting()
    assert [a["file"] for a in again["already"]] == ["copy.csv"]


def test_the_statements_folder_goes_in_the_documents_folder_windows_shows(tmp_path, monkeypatch):
    # OneDrive often moves Documents and leaves an empty C:\Users\name\Documents behind
    home, moved = tmp_path / "home", tmp_path / "home" / "OneDrive" / "Documents"
    (home / "Documents").mkdir(parents=True)
    moved.mkdir(parents=True)
    monkeypatch.setattr(inbox.pathlib.Path, "home", classmethod(lambda cls: home))
    monkeypatch.setattr(inbox, "_windows_documents", lambda: moved)
    assert inbox.default_folder().parent == moved
    monkeypatch.setattr(inbox, "_windows_documents", lambda: None)   # not Windows, or no answer
    assert inbox.default_folder().parent == home / "Documents"


def test_the_reminder_repeats_monthly_and_last_day_means_the_last_day():
    config.settings.update({"reminder_day": 31})
    ics = reminders.ics(dt.date(2026, 2, 3))
    assert "RRULE:FREQ=MONTHLY;BYMONTHDAY=-1" in ics and ics.endswith("\r\n")
    assert reminders.next_due(dt.date(2026, 2, 3)) == dt.date(2026, 2, 28)


# ============================================================================ money in other currencies

def test_rand_is_converted_and_never_counted_one_for_one():
    with db.tx() as c:
        c.execute("INSERT INTO platforms(name,provider,wrapper,cash,currency) VALUES('Capitec','bank','cash',10000,'ZAR')")
    bs = sections.balance_sheet()
    assert bs["unconverted"] == ["Capitec"] and "Cash" not in bs["own"]     # no rate yet: left out, and said so
    from .test_whiskers import price_series
    price_series("FX:ZAR", [23.26] * 5)                       # rand per pound
    from folio.market import store
    store._fx_cache.clear()
    bs = sections.balance_sheet()
    assert bs["own"]["Cash"] == pytest.approx(430, abs=1) and not bs["unconverted"]


# ============================================================================ market data that stops

def test_ons_monthly_series_are_read_and_odd_rows_skipped():
    from folio.market import sources
    payload = {"description": {"cdid": "D7BT"}, "years": [{"date": "2025", "value": "138.2"}],
               "months": [{"date": "2026 JUL", "value": "141.0"}, {"date": "2026 AUG", "value": "141.4"},
                          {"date": "2026 SEP", "value": ""}, {"date": "rubbish", "value": "1"}]}
    assert sources.parse_ons_json(payload) == [("2026-07-01", 141.0), ("2026-08-01", 141.4)]
    with pytest.raises(sources.SourceError):
        sources.parse_ons_json({"years": []})


def test_a_light_whose_source_stopped_publishing_is_left_out_not_shown_as_current():
    from folio.engine import signals, themes
    from folio.market import store
    defn = next(d for d in themes.lights_for("inflation") if d["key"] == "uk_cpi")
    key = store.ref_key("UKCPI")

    def monthly(end: dt.date, n: int = 30):
        out, y, m = [], end.year, end.month
        for i in range(n):
            out.append((f"{y:04d}-{m:02d}-01", 100.0 + (n - i) * 0.3))
            y, m = (y - 1, 12) if m == 1 else (y, m - 1)
        return sorted(out)

    store._save(key, {"points": monthly(dt.date(2025, 3, 1))}, "fred:GBRCPIALLMINMEI", replace=True)
    r = signals.reading(defn)
    assert r["level"] == "none" and "1 Mar 2025" in r["words"]
    # a figure from last month is what a monthly series normally looks like
    last_month = (dt.date.today().replace(day=1) - dt.timedelta(days=1)).replace(day=1)
    store._save(key, {"points": monthly(last_month)}, "ons:test", replace=True)
    assert signals.reading(defn)["level"] in ("green", "amber", "red")


def test_the_app_asks_before_moving_the_main_watch_and_never_when_it_was_picked():
    from folio.engine import portfolio
    pid = add_platform(cash=0.0)
    add_position(pid, "SGLN", exchange="LSE", name="iShares Physical Gold ETC", broker_value=6000.0, valued_at=db.now())
    add_position(pid, "IGLT", exchange="LSE", name="iShares Core UK Gilts UCITS ETF", broker_value=3000.0,
                 valued_at=db.now())
    pos = portfolio.positions()
    tot = portfolio.totals(pos)
    assert rules._watch_moved(pos, tot) is None                  # the original AI watch is left alone
    themes.set_watch("rates", ["rates", "gold"], None)
    db.set_meta("builder", {"v": 2, "answers": {"watch_main": "auto"}, "at": {}, "mood": {}})
    w = rules._watch_moved(pos, tot)
    assert w and w["level"] == "info" and "Make gold your main watch?" in w["detail"]
    db.set_meta("builder", {"v": 2, "answers": {"watch_main": "rates"}, "at": {}, "mood": {}})
    assert rules._watch_moved(pos, tot) is None                  # picked by name: theirs to keep
    db.set_meta("builder", {"v": 2, "answers": {"views": {"rates": {"action": "parachute"}}}, "at": {}, "mood": {}})
    assert rules._watch_moved(pos, tot) is None                  # watched for its parachute, whatever it carries


def test_the_mortgage_line_says_whether_it_outlives_work():
    retirement.save({"dob": _years_ago(55), "stop_age": 60})
    payoff = dt.date.today() + dt.timedelta(days=round(8 * 365.25))
    assert sections._clear_by(payoff).startswith("clears ") and "after" not in sections._clear_by(payoff)
    _on(retirement=True)                                        # the stopping age is the Retirement section's
    assert sections._clear_by(payoff).endswith("3 years after you stop work")
    assert sections._clear_by(dt.date.today() + dt.timedelta(days=365)).endswith("before you stop work")


def test_a_theme_outside_the_watch_that_turns_red_is_listed_but_never_moves_the_dial():
    from folio.engine import signals
    from folio.market import store
    from .test_whiskers import price_series
    price_series(store.ref_key("BRENT"), [100.0] * 260 + [100.0 - i for i in range(1, 46)], currency="USD")
    moved = signals.moved_outside()
    oil = next((x for x in moved if x["theme"] == "oil"), None)
    assert oil and any("fall" in r["title"].lower() for r in oil["reds"])
    assert all(x["theme"] != "ai" for x in moved)              # the original watch is watched
    assert "oil" not in signals.evaluate()["groups"]          # and the dial never counts it


def test_a_thesis_keeps_its_factor_grades_and_the_brief_shows_them():
    from folio.web import routes2
    pid = add_platform()
    iid = add_position(pid, "AZN", exchange="LSE", name="AstraZeneca", broker_value=5000.0, valued_at=db.now())
    routes2.api_save_thesis(None, {"iid": str(iid)}, {}, {"why": "Oncology pipeline", "breakers": "Pipeline fails",
                            "factors": {"valuation": "C-", "growth": "A", "momentum": "Z", "nonsense": "A"}})
    row = db.one("SELECT factors FROM theses WHERE instrument_id=?", (iid,))
    f = json.loads(row["factors"])
    assert f["valuation"] == "C-" and f["growth"] == "A" and "momentum" not in f and "nonsense" not in f and f["on"]
    brief = builder.brief()
    assert "Factor grades" in brief and "Oncology pipeline" in brief and "Pipeline fails" in brief
    routes2.api_save_thesis(None, {"iid": str(iid)}, {}, {"why": "Oncology pipeline", "factors": {}})
    assert db.one("SELECT factors FROM theses WHERE instrument_id=?", (iid,))["factors"] is None


# ============================================================================ the sample

def test_the_sample_comes_out_completely_and_keeps_your_own_details():
    sample.load()
    assert db.scalar("SELECT COUNT(*) FROM pensions", (), 0) == 2 and mort.all_mortgages()
    sample.remove()
    for t in ("pensions", "mortgages", "assets", "platforms"):
        assert not db.scalar(f"SELECT COUNT(*) FROM {t}", (), 0), t  # nosec B608
    assert retirement.settings()["dob"] is None
    # somebody's own date of birth is never replaced by the sample's
    retirement.save({"dob": "1964-02-02"})
    sample.load()
    sample.remove()
    assert retirement.settings()["dob"] == "1964-02-02"


# ============================================================================ the local server

@pytest.fixture()
def server():
    from folio.web import server as srv
    httpd, port = srv.serve(port=srv._free_port(18866), block=False)
    yield port
    httpd.shutdown()


def _get(port, path):
    try:
        with urllib.request.urlopen(f"http://127.0.0.1:{port}{path}", timeout=60) as r:
            return r.status, r.read(), r.headers.get("Content-Type", "")
    except urllib.error.HTTPError as e:
        return e.code, e.read(), ""


NEW_SCREENS = ("/api/assets", "/api/mortgages", "/api/pensions", "/api/retirement", "/api/overview",
               "/api/builder", "/api/builder/draft", "/api/builder/draft?mode=fresh", "/api/builder/written",
               "/api/xray", "/api/watch", "/api/ai", "/api/allocation", "/api/institutions?q=hargreaves",
               "/api/inbox", "/api/reminder")


def test_every_new_screen_answers_on_an_empty_copy(server):
    for path in NEW_SCREENS:
        status, body, ctype = _get(server, path)
        assert status == 200, (path, body[:200])
        assert "error" not in json.loads(body), path


def test_every_new_screen_answers_with_the_sample_household(server):
    sample.load()
    for path in NEW_SCREENS + ("/api/mortgages/1", "/api/instrument/1/themes"):
        status, body, _ = _get(server, path)
        assert status == 200, (path, body[:300])
    status, body, ctype = _get(server, "/api/builder/written.html")
    assert status == 200 and b"<h1>" in body and "text/html" in ctype
    status, body, ctype = _get(server, "/api/builder/brief")
    assert status == 200 and body.startswith(b"# Plan brief for Claude") and "markdown" in ctype


def test_text_size_is_a_setting_with_three_sizes(server):
    def patch(v):
        req = urllib.request.Request(f"http://127.0.0.1:{server}/api/settings", method="PATCH",
                                     data=json.dumps({"text_size": v}).encode(), headers={"Content-Type": "application/json"})
        try:
            with urllib.request.urlopen(req, timeout=30) as r:
                return r.status
        except urllib.error.HTTPError as e:
            return e.code
    assert patch("largest") == 200 and config.settings.get("text_size") == "largest"
    assert patch("enormous") == 400 and config.settings.get("text_size") == "largest"
    status, body, _ = _get(server, "/api/bootstrap")
    assert json.loads(body)["settings"]["text_size"] == "largest"


def test_the_new_routes_never_use_error_for_anything_but_failure():
    from .test_whiskers import ROOT
    src = (ROOT / "folio" / "web" / "routes2.py").read_text(encoding="utf-8")
    assert '"error":' not in src


def test_a_built_app_carries_the_institution_catalogue_and_the_test_copy_keeps_apart():
    from .test_whiskers import ROOT
    spec = (ROOT / "build" / "whiskers.spec").read_text(encoding="utf-8")
    assert '"institutions/data"' in spec and "folio.web.routes2" in spec
    assert "folio.budget" not in spec and "phonecopy" not in spec
    bat = (ROOT / "Try Whiskers 1.1 (test copy).bat").read_text(encoding="utf-8")
    assert 'WHISKERS_DATA_DIR=%LOCALAPPDATA%\\Whiskers 1.1 test' in bat


# ============================================================================ the optional sections

def test_sections_off_means_nothing_shown_and_nothing_counted():
    sample.load()                                       # brings a home, a mortgage and two pensions
    from folio.engine import export, portfolio
    assert sections.on() == {"mortgage": False, "retirement": False}
    assert sections.overview() is None
    bs = sections.balance_sheet()
    assert "Home and property" not in bs["own"] and "Pensions" not in bs["own"] and not bs["owe"].get("Mortgage")
    assert retirement.gather()["mortgages"] == []       # the retirement plan ignores a mortgage that is off
    from openpyxl import load_workbook
    import io
    names = load_workbook(io.BytesIO(export.build())).sheetnames
    assert not {"Mortgage", "Pensions", "Retirement", "Net worth"} & set(names)
    total_before = portfolio.totals()["total"]
    _on(mortgage=True, retirement=True)
    assert portfolio.totals()["total"] == pytest.approx(total_before)     # investments never change
    ov = sections.overview()
    assert [x["key"] for x in ov["lines"]] == ["mortgage", "retirement"]
    assert ov["own"]["Home and property"] > 0 and ov["owe"]["Mortgage"] > 0 and ov["own"]["Pensions"] > 0
    names = load_workbook(io.BytesIO(export.build())).sheetnames
    assert {"Mortgage", "Pensions"} <= set(names)


def test_one_section_counts_only_its_own_parts():
    sample.load()
    _on(mortgage=True)
    bs = sections.balance_sheet()
    assert bs["owe"]["Mortgage"] > 0 and "Pensions" not in bs["own"]
    assert [x["key"] for x in sections.overview()["lines"]] == ["mortgage"]


def test_the_sections_that_suit_someone_are_offered_and_no_more():
    assert sections.advice(36, False, 25, False) == {}                       # Allan: nothing
    mum = sections.advice(69, True, 0, True)
    assert set(mum) == {"mortgage", "retirement"}
    assert "retirement" in sections.advice(58, False, 8, False) and "mortgage" not in sections.advice(58, False, 8, False)


# ============================================================================ the plan builder

ALLAN = {"whiskers_plan": 1, "replace_sleeves": True, "sleeves": [
    {"name": "Global core index", "target": 23}, {"name": "Technology & AI", "target": 23, "ai_share": 100},
    {"name": "Healthcare", "target": 9}, {"name": "Industrials & Materials", "target": 7.5},
    {"name": "Commodities", "target": 7}, {"name": "Real assets", "target": 6.5}, {"name": "Consumer", "target": 5.5},
    {"name": "Financials", "target": 4.5}, {"name": "Cash", "target": 4, "is_cash": True}, {"name": "Gilts", "target": 3},
    {"name": "Energy", "target": 3}, {"name": "Defence", "target": 2}, {"name": "Speculative", "target": 2}],
    "rules": [{"kind": "company_cap", "subject": "Alphabet", "trigger": 9.3, "target": 8, "basis": "direct"}]}
ALLAN_ANSWERS = {
    "need_when": "over20", "emergency": "yes", "income_pref": "growth", "age_band": "under40", "home": "no",
    "fall": "buy", "did": "bought", "experience": "few", "prefer": "growth", "hands": "enjoy", "time": "lots",
    "worries": ["bubble", "rates", "inflation"], "avoid": ["none"],
    "views": {"ai": {"believe": 5, "worry": 4, "action": "parachute", "why": "The biggest change in my lifetime.",
                     "change": "Hyperscalers cutting capex two quarters running."},
              "healthcare": {"believe": 4, "worry": 2, "action": "ride"},
              "industrials": {"believe": 4, "worry": 2, "action": "ride"},
              "metals": {"believe": 3, "worry": 2, "action": "ride"},
              "property": {"believe": 3, "worry": 3, "action": "ride"},
              "consumer": {"believe": 3, "worry": 1, "action": "ride"},
              "financials": {"believe": 3, "worry": 2, "action": "ride"},
              "oil": {"believe": 2, "worry": 2, "action": "little"},
              "defence": {"believe": 3, "worry": 2, "action": "little"},
              "speculative": {"believe": 3, "worry": 4, "action": "little"}}}


def _targets():
    return {r["name"]: r["target"] for r in db.rows("SELECT name, target FROM sleeves")}


def test_risk_is_the_lower_of_willing_and_able_and_a_real_sale_counts_most():
    bold = builder.risk_profile({"fall": "buy", "experience": "long", "prefer": "growth", "need_when": "over20"}, {})
    timid = builder.risk_profile({"fall": "sell", "experience": "new", "prefer": "steady", "need_when": "over20"}, {})
    assert timid["growth"] < bold["growth"] and timid["growth"] == min(timid["willing"], timid["able"])
    sold = builder.risk_profile({"fall": "buy", "did": "sold", "experience": "long", "prefer": "growth"}, {})
    assert sold["willing"] < bold["willing"]
    soon = builder.risk_profile({"fall": "buy", "experience": "long", "prefer": "growth", "need_when": "under5"}, {})
    assert soon["growth"] < bold["growth"]


def test_the_builder_asks_nothing_the_retirement_section_already_knows():
    retirement.save({"dob": _years_ago(70.5), "stop_age": 72, "spend_essential": 1500})
    a = builder.answers()
    assert "need_when" not in a                          # the section is off: asked in the builder
    _on(retirement=True)
    a = builder.answers()
    assert a["need_when"] == "under5" and a["age_band"] == "65to74"
    builder.save_answers({"need_when": "over20"})        # a derived answer can't be overruled from here
    assert builder.answers()["need_when"] == "under5"


def test_answers_are_saved_as_given_and_views_merge():
    builder.save_answers({"views": {"ai": {"believe": 5}}})
    builder.save_answers({"views": {"ai": {"action": "parachute", "why": "  Real returns.  "}}})
    builder.save_answers({"views": {"nonsense": {"believe": 3}, "gold": {"believe": 9, "action": "moon"}}})
    v = builder.answers()["views"]
    assert v["ai"] == {"believe": 5, "action": "parachute", "why": "Real returns."}
    assert "nonsense" not in v and v["gold"] == {"believe": 5}
    builder.save_answers({"fall": "panic"})              # not one of the options: ignored
    assert "fall" not in builder.answers()


def test_a_personal_plan_is_reviewed_not_replaced():
    sample.load()
    plan.import_plan(ALLAN)
    builder.save_answers(ALLAN_ANSWERS)
    before = _targets()
    d = builder.draft()
    assert d["mode"] == "review"
    sugg = d["review"]["suggestions"]
    assert not [s for s in sugg if s["kind"] in ("target", "add", "move")], [s["title"] for s in sugg]
    kinds = {s["change"].get("rule", {}).get("kind") for s in sugg if s["kind"] == "rule"}
    assert {"ai_cap", "theme_step", "any_company_cap"} <= kinds          # the parachute he asked for
    step = next(s for s in sugg if s["kind"] == "rule" and s["change"]["rule"]["kind"] == "theme_step")
    assert step["change"]["rule"]["into"] == "Cash"      # rising rates worry him: cash before gilts
    watch = next(s for s in sugg if s["kind"] == "watch")
    assert watch["ticked"] is False                      # the AI dial stays; extra lights are his choice
    assert d["sections"] == []
    out = builder.apply({"mode": "review", "ticked": [s["id"] for s in sugg if s["ticked"]]})
    assert _targets() == before and out["ok"]
    assert db.scalar("SELECT COUNT(*) FROM rules WHERE kind='theme_step' AND subject='ai'") == 1
    assert themes.watch()["legacy"] is True              # unticked, so untouched
    why = db.one("SELECT why, breakers FROM sleeves WHERE name='Technology & AI'")
    assert why["why"] == "The biggest change in my lifetime." and "capex" in why["breakers"]


def test_answers_that_contradict_a_plan_are_suggested_but_left_unticked():
    sample.load()
    plan.import_plan(ALLAN)
    builder.save_answers(dict(ALLAN_ANSWERS, avoid=["weapons"]))
    sugg = builder.draft()["review"]["suggestions"]
    wind = next(s for s in sugg if s["id"].startswith("avoid:"))
    assert wind["ticked"] is False and wind["change"]["sleeve"] == "Defence" and wind["change"]["balance"] == "Global core index"
    builder.apply({"mode": "review", "ticked": [wind["id"]]})
    t = _targets()
    assert t["Defence"] == 0 and t["Global core index"] == 25 and sum(t.values()) == pytest.approx(100)


def test_a_fresh_plan_follows_the_answers_and_changes_nothing_until_used():
    sample.load()
    before = _targets()
    builder.save_answers({"need_when": "under5", "draw_rate": "4", "fall": "worry", "experience": "few",
                          "prefer": "balanced", "hands": "funds", "worries": ["inflation", "rates"],
                          "views": {"healthcare": {"believe": 4, "action": "ride"}, "oil": {"believe": 3, "action": "dip"}}})
    d = builder.draft("fresh")
    sl = {s["name"]: s for s in d["plan"]["sleeves"]}
    assert sum(s["target"] for s in sl.values()) == pytest.approx(100)
    assert sl["Global core index"]["target"] > 30 and sl["Healthcare"]["target"] <= 6     # funds only: a big core
    assert "Index-linked gilts" in sl and "Short-dated gilts" in sl                     # the worries' defences
    assert any(r["kind"] == "theme_dip" and r["subject"] == "oil" for r in d["plan"]["rules"])
    assert _targets() == before
    builder.apply({"mode": "fresh"})
    assert "Global core index" in _targets() and "Short-dated gilts" in _targets()
    assert db.scalar("SELECT COUNT(*) FROM rules WHERE kind='theme_dip'") == 1


def test_someone_near_seventy_with_a_mortgage_is_offered_both_sections():
    builder.save_answers({"age_band": "65to74", "home": "yes", "need_when": "now"})
    d = builder.draft()
    assert {x["key"] for x in d["sections"]} == {"mortgage", "retirement"}
    builder.apply({"mode": "fresh", "parts": {"sleeves": False, "rules": False, "watch": False, "plans": False},
                   "sections": {"mortgage": True, "retirement": True}})
    assert sections.on() == {"mortgage": True, "retirement": True}


def test_the_written_plan_carries_the_reasons_for_each_theme():
    builder.save_answers({"views": {"gold": {"believe": 4, "worry": 2, "action": "ride", "why": "Central banks buy it.",
                                             "change": "Real yields above 3%."}}})
    w = builder.written_plan()
    cards = next(s for s in w["sections"] if s["kind"] == "cards")["lines"]
    assert cards[0]["title"] == "Gold" and cards[0]["why"] == "Central banks buy it." and cards[0]["action"] == "Ride it"


# ============================================================================ the parachute

def _lights_turn(ref: str, fall: float):
    from folio.market import store
    from .test_whiskers import price_series
    price_series(store.ref_key(ref), [100.0 + i * 0.1 for i in range(300)] + [130.0 * (1 - fall)] * 30, currency="USD")


def test_a_parachute_step_shows_when_its_lights_turn_and_rests_once_taken():
    from folio.engine import portfolio
    pid = add_platform(cash=1000.0)
    add_position(pid, "SGLN", exchange="LSE", name="iShares Physical Gold ETC", broker_value=9000.0, valued_at=db.now())
    plan.import_plan({"whiskers_plan": 2, "sleeves": [{"name": "Gold", "target": 90, "theme": "gold"},
                                                      {"name": "Cash", "target": 10, "is_cash": True}],
                      "rules": [{"kind": "theme_step", "subject": "gold", "trigger": 2, "target": 33, "sleeve": "Gold",
                                 "into": "Cash", "note": "Parachute: when 2 of gold's 3 lights are red, move a third into Cash."}]})
    db.execute("UPDATE instruments SET sleeve_id=(SELECT id FROM sleeves WHERE name='Gold')")
    pos = portfolio.positions()
    tot = portfolio.totals(pos)
    calm = [s for s in rules.rule_status(pos, tot) if s["rule"]["kind"] == "theme_step"][0]
    assert calm["level"] in ("none", "green")
    _lights_turn("GOLD", 0.35)                                # a bear market and a broken trend
    st = [s for s in rules.rule_status(pos, tot) if s["rule"]["kind"] == "theme_step"][0]
    assert st["level"] == "red" and st["now"] >= 2 and st["move"] == pytest.approx(3000, rel=0.05)
    warns = [w for w in rules.evaluate(pos, tot, {"sleeves": [], "targets_add_up": True}) if w["kind"] == "step"]
    assert warns and "the step you agreed applies" in warns[0]["title"] and "Cash" in warns[0]["detail"]
    rules.mark_done(st["rule"]["id"])
    assert not [w for w in rules.evaluate(pos, tot, {"sleeves": [], "targets_add_up": True}) if w["kind"] == "step"]
    _lights_turn("GOLD", -0.05)                               # recovered: the step re-arms
    rules.rule_status(pos, tot)
    assert str(st["rule"]["id"]) not in rules.steps_done()


def test_a_buying_point_shows_after_the_fall_the_person_was_waiting_for():
    from folio.engine import portfolio
    pid = add_platform(cash=5000.0)
    add_position(pid, "SHEL", exchange="LSE", name="Shell", broker_value=5000.0, valued_at=db.now())
    plan.import_plan({"whiskers_plan": 2, "sleeves": [{"name": "Energy", "target": 50, "theme": "oil"},
                                                      {"name": "Cash", "target": 50, "is_cash": True}],
                      "rules": [{"kind": "theme_dip", "subject": "oil", "trigger": 30, "target": 5, "sleeve": "Cash",
                                 "into": "Energy"}]})
    pos = portfolio.positions()
    tot = portfolio.totals(pos)
    _lights_turn("ENERGY", 0.40)
    st = [s for s in rules.rule_status(pos, tot) if s["rule"]["kind"] == "theme_dip"][0]
    assert st["level"] == "amber" and st["now"] >= 30 and st["move"] == pytest.approx(500)


def test_any_one_company_is_trimmed_once_it_passes_the_limit():
    from folio.engine import portfolio
    pid = add_platform(cash=0.0)
    add_position(pid, "NVDA", name="NVIDIA", broker_value=1500.0, valued_at=db.now())
    add_position(pid, "VWRP", exchange="LSE", name="Vanguard FTSE All-World (Acc)", broker_value=8500.0, valued_at=db.now())
    plan.import_plan({"whiskers_plan": 2, "rules": [{"kind": "any_company_cap", "trigger": 10, "target": 8, "basis": "direct"}]})
    pos = portfolio.positions()
    tot = portfolio.totals(pos)
    warns = [w for w in rules.evaluate(pos, tot, {"sleeves": [], "targets_add_up": True}) if w["kind"] == "any_company_cap"]
    assert len(warns) == 1 and warns[0]["level"] == "red" and "NVIDIA" in warns[0]["title"]
    assert "£700" in warns[0]["detail"]                       # 15% back to 8% of £10,000


def test_plan_files_carry_reasons_and_steps_both_ways():
    plan.import_plan({"whiskers_plan": 2, "replace_sleeves": True, "sleeves": [
        {"name": "Core", "target": 70}, {"name": "Technology & AI", "target": 25, "theme": "ai", "why": "My reason.",
                                          "change": "My breaker."}, {"name": "Cash", "target": 5, "is_cash": True}],
        "rules": [{"kind": "theme_step", "subject": "ai", "trigger": 2, "target": 33, "sleeve": "Technology & AI", "into": "Cash"}]})
    out = plan.export_plan()
    tech = next(s for s in out["sleeves"] if s["name"] == "Technology & AI")
    assert tech["why"] == "My reason." and tech["change"] == "My breaker." and tech["theme"] == "ai"
    step = next(r for r in out["rules"] if r["kind"] == "theme_step")
    assert step["sleeve"] == "Technology & AI" and step["into"] == "Cash"
    with pytest.raises(plan.PlanError):
        plan.import_plan({"whiskers_plan": 2, "rules": [{"kind": "theme_step", "subject": "ai", "trigger": 2,
                                                         "sleeve": "Technology & AI", "into": "Nowhere"}]})
    plan.import_plan({"whiskers_plan": 2, "sleeves": [{"name": "Core", "target": 70}]})
    assert db.one("SELECT why FROM sleeves WHERE name='Technology & AI'")["why"] == "My reason."   # kept


# ============================================================================ the brief and the file that comes back

def test_the_brief_has_what_claude_needs_and_amounts_only_when_asked():
    sample.load()
    builder.save_answers(ALLAN_ANSWERS)
    b = builder.brief()
    for part in ("## What to send back", "### The plan file", "## Their answers", "## Their views, in their own words",
                 "## What they hold", "## Their current plan", "## The X-Ray", "## Theme keys"):
        assert part in b, part
    assert "The biggest change in my lifetime." in b and "Ride it, with a parachute" in b
    assert "£" not in b.split("## What they hold")[1].split("## Their current plan")[0]
    assert "£" in builder.brief(amounts=True).split("## What they hold")[1].split("## Their current plan")[0]


def test_a_returned_plan_is_checked_before_anything_changes():
    sample.load()
    builder.save_answers(ALLAN_ANSWERS)
    before = _targets()
    bad = builder.check_plan(b'{"whiskers_plan": 2, "replace_sleeves": true, "sleeves": [{"name": "A", "target": 60}]}')
    assert not bad["ok"] and any("60%" in e for e in bad["errors"]) and bad["token"] is None
    assert builder.check_plan(b"not a plan at all")["ok"] is False
    reply = ("Here is your plan:\n```json\n" + json.dumps(dict(ALLAN, whiskers_plan=2)) + "\n```\nAnd why.").encode()
    good = builder.check_plan(reply)
    assert good["ok"] and good["token"]
    assert any("parachute" in w for w in good["warnings"])       # he asked for a parachute; the file has none
    assert any(x["change"] == "new" for x in good["diff"]["sleeves"])
    assert _targets() == before                                  # checking changes nothing
    with pytest.raises(ValueError):
        builder.use_checked("not-the-token")
    builder.use_checked(good["token"])
    assert _targets()["Technology & AI"] == 23 and sum(_targets().values()) == pytest.approx(100)


def test_a_step_taken_can_move_the_targets_too(server):
    plan.import_plan({"whiskers_plan": 2, "replace_sleeves": True, "sleeves": [
        {"name": "Technology & AI", "target": 30}, {"name": "Core", "target": 60}, {"name": "Cash", "target": 10, "is_cash": True}],
        "rules": [{"kind": "theme_step", "subject": "ai", "trigger": 2, "target": 33, "sleeve": "Technology & AI", "into": "Cash"}]})
    rid = db.scalar("SELECT id FROM rules WHERE kind='theme_step'")
    req = urllib.request.Request(f"http://127.0.0.1:{server}/api/rules/{rid}/done", method="POST",
                                 data=json.dumps({"targets": True}).encode(), headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=30) as r:
        out = json.loads(r.read())
    assert out["moved"] == {"points": 10.0, "from": "Technology & AI", "to": "Cash"}
    assert _targets() == {"Technology & AI": 20, "Core": 60, "Cash": 20}
    assert str(rid) in rules.steps_done()
